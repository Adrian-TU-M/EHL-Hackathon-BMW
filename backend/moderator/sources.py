"""Read only the nine supplied artifacts. Raw text and derived catalogs stay private."""

from __future__ import annotations

import hashlib
import json
import re
from collections import Counter
from pathlib import Path

import openpyxl
import pymupdf as fitz

from .models import Catalog, Evidence, Source

FEEDBACK = {"F70": (4076, 2128), "G60": (5005, 4000), "G70": (4690, 4000)}
EU = {
    "DE",
    "FR",
    "SE",
    "FI",
    "DK",
    "LV",
    "ES",
    "IT",
    "NL",
    "BE",
    "AT",
    "PL",
    "PT",
    "IE",
    "CZ",
    "GR",
    "HU",
    "RO",
    "BG",
    "HR",
    "SK",
    "SI",
    "EE",
    "LT",
    "LU",
    "MT",
    "CY",
}
EUROPE = EU | {"GB", "NO", "CH", "IS"}
POWERTRAINS = {"BEVE": "Electric", "PHEV": "Plug-in hybrid", "ICE": "Combustion"}
DRIVE_IDS = {
    "F70_feedback": "1IFcCXAJ4TGSQM8RYdoggGTOEp2KDrhXC",
    "G60_feedback": "1YuCFSiGYza-60f3lrMACtWzEuxIt9VJg",
    "G70_feedback": "1uaObSrfXDcDzXRiNVtQHAy9RCDcaZaSd",
    "studies": "1wb2utLErwdrIVXIw_MN8o0zQ3H1VlTid",
    "sales": "1Etq2cZ5aq7LAtKuT5Kno2uOKpyBecFNy",
}
# These are candidate extraction rules, not a fixed agenda. Unknown classifications
# remain independent candidates; original text is always inspected as well.
RULES = {
    "climate": (
        "Climate usability",
        "Interior",
        "How should cabin climate remain easy to adjust?",
        r"climate|air.condition|ventilation|heating|temperature|a/c|defrost",
    ),
    "app": (
        "App reliability",
        "Displays",
        "How should the companion app remain reliable?",
        r"my bmw|mybmw|\bapp\b|smartphone application",
    ),
    "services": (
        "Connected services",
        "Displays",
        "Which connected services should remain dependable and understandable?",
        r"connecteddrive|connected service|subscription|online service|digital service",
    ),
    "access": (
        "Keys and access",
        "Exterior",
        "How should access and vehicle locking work reliably?",
        r"digital key|\bkeys?\b|unlock|locking|door handle|comfort access",
    ),
    "charging": (
        "Charging experience",
        "Charging",
        "How should charging be simpler and more dependable?",
        r"charg|plug.and.charge|bidirectional",
    ),
    "range": (
        "Range and efficiency",
        "Charging",
        "How should usable range meet journey needs?",
        r"\brange\b|consumption|efficiency|battery capacity",
    ),
    "luggage": (
        "Luggage usability",
        "Trunk",
        "How should luggage space work for everyday loads?",
        r"luggage|\btrunk\b|boot space|cargo|tailgate",
    ),
    "noise": (
        "Noise and quality",
        "Interior",
        "How should noise and perceived quality be improved while preserving comfort?",
        r"noise|rattle|squeak|build quality|workmanship|acoustic",
    ),
    "controls": (
        "Essential controls",
        "Driving controls",
        "How should essential controls remain accessible during driving?",
        r"control|button|switch|operat|idrive|panoramic",
    ),
    "display": (
        "Display and information",
        "Displays",
        "How should driving information remain legible and useful?",
        r"display|screen|head.up|instrument|navigation|infotainment",
    ),
    "seats": (
        "Seating comfort",
        "Seats",
        "How should seats support comfort and adjustment?",
        r"seat|seating",
    ),
    "assistance": (
        "Driver assistance",
        "Driving controls",
        "How should assistance support drivers without unwanted intervention?",
        r"assistan|lane|cruise|parking|autonom|warning",
    ),
    "ride": (
        "Ride and handling",
        "Driving controls",
        "How should ride and handling balance comfort and control?",
        r"suspension|steering|handling|ride comfort|braking|brakes",
    ),
    "space": (
        "Cabin usability",
        "Interior",
        "How should cabin space and storage support everyday use?",
        r"interior|cabin|storage|legroom|cup.holder",
    ),
    "exterior": (
        "Exterior usability",
        "Exterior",
        "How should exterior design balance usability and valued appearance?",
        r"exterior|styling|design|light|mirror|wiper",
    ),
}


def topics_for(text: str, classifications: list[str] | None = None) -> list[str]:
    found = [key for key, rule in RULES.items() if re.search(rule[3], text, re.I)]
    for category in classifications or []:
        if not category or category == "no_class_found":
            continue
        matches = [
            key for key, rule in RULES.items() if re.search(rule[3], category, re.I)
        ]
        found.extend(matches)
    if not found:
        meaningful = [
            c
            for c in classifications or []
            if c
            and c != "no_class_found"
            and c.lower()
            not in {"not user friendly", "difficult to use", "defect", "wants", "likes"}
        ]
        if meaningful:
            found.append(
                "category-"
                + re.sub(r"[^a-z0-9]+", "-", meaningful[-1].lower()).strip("-")
            )
    return sorted(set(found or ["uncategorized"]))


def load_catalog(folder: Path, cache: Path | None = None) -> Catalog:
    names = [f"{v}_feedback_hackathon.xlsx" for v in FEEDBACK] + [
        "F70_G60_G68_G70_customer_studies.xlsx",
        "sales_volumes.xlsx",
        *[f"{v}_OptionList.PDF" for v in FEEDBACK],
        "BMW-Group-Report-2025-en.pdf",
    ]
    # Accept the supplied BMWHack root, its BMW-Data child, or the committed
    # private-data/inputs layout without copying or modifying source documents.
    files = {
        name: next((candidate for candidate in
                    (folder / name, folder / "BMW-Data" / name, folder / "inputs" / name, folder.parent / name)
                    if candidate.is_file()), folder / name)
        for name in names
    }
    missing = [n for n in names if not files[n].is_file()]
    if missing:
        raise FileNotFoundError("Missing supplied artifacts: " + ", ".join(missing))
    hashes = {n: hashlib.sha256(files[n].read_bytes()).hexdigest() for n in names}
    snapshot = hashlib.sha256(
        json.dumps({"extractor": "v6", **hashes}, sort_keys=True).encode()
    ).hexdigest()
    target = cache / f"{snapshot}.json" if cache else None
    if target and target.exists():
        return Catalog.model_validate_json(target.read_text())
    sources, evidence, warnings = [], [], []

    def source(name, sid, domain, **kwargs):
        s = Source(id=sid, name=name, domain=domain, sha256=hashes[name], **kwargs)
        sources.append(s)
        return s

    for vehicle, expected in FEEDBACK.items():
        name = f"{vehicle}_feedback_hackathon.xlsx"
        sid = f"{vehicle}_feedback"
        s = source(
            name,
            sid,
            "customer",
            source_url=f"https://docs.google.com/spreadsheets/d/{DRIVE_IDS[sid]}/edit",
            limitations=[
                "Feedback IDs identify observations, not unique customers.",
                "Dates are not supplied; populations are not representative samples.",
            ],
        )
        wb = openpyxl.load_workbook(files[name], read_only=True, data_only=True)
        ws = wb["Feedback_Explorer"]
        observations = {}
        count = 0
        for rownum, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            if row[0] is None:
                continue
            count += 1
            fid = str(row[0])
            if fid not in observations:
                observations[fid] = Evidence(
                    id=f"{sid}:{fid}",
                    source_id=sid,
                    domain="customer",
                    location="",
                    passage=str(row[8] or ""),
                    vehicle=vehicle,
                    country=str(row[4] or "") or None,
                    powertrain=POWERTRAINS.get(str(row[6]), str(row[6] or "unknown")),
                    kind="feedback",
                    observation_id=fid,
                    limitations=[
                        "Observation date and unique-customer identity are unknown."
                    ],
                )
            e = observations[fid]
            if e.passage != str(row[8] or "") or e.country != (
                str(row[4] or "") or None
            ):
                raise ValueError(
                    f"Conflicting text/scope for {e.id}; ingestion needs review"
                )
            e.classification_rows.append(rownum)
            e.classifications.extend(
                str(c) for c in row[9:12] if c and str(c) not in e.classifications
            )
            if row[7] and str(row[7]) not in e.feedback_types:
                e.feedback_types.append(str(row[7]))
        for e in observations.values():
            e.location = "Feedback_Explorer!" + ",".join(
                f"A{r}:L{r}" for r in e.classification_rows
            )
            e.topics = topics_for(e.passage, e.classifications)
        s.coverage = {
            "classification_rows": count,
            "distinct_feedback_ids": len(observations),
            "countries": dict(Counter(e.country for e in observations.values())),
            "unclassified_observations": sum(
                not any(c != "no_class_found" for c in e.classifications)
                for e in observations.values()
            ),
        }
        if (count, len(observations)) != expected:
            raise ValueError(
                f"Source count mismatch for {sid}: {(count, len(observations))}, expected {expected}"
            )
        evidence.extend(observations.values())
        wb.close()

    name = "F70_G60_G68_G70_customer_studies.xlsx"
    s = source(
        name,
        "studies",
        "customer",
        source_url=f"https://docs.google.com/spreadsheets/d/{DRIVE_IDS['studies']}/edit",
        limitations=[
            "Respondent counts are unknown.",
            "Mean rating scale is undocumented.",
            "US sample total 1 is a normalized column distribution.",
            "China G68 is not G60.",
            "EU does not mean all Europe.",
        ],
    )
    wb = openpyxl.load_workbook(files[name], data_only=True)
    ws = wb["CN_EU_2025"]
    for col, vehicle, region in [
        (2, "G68", "China"),
        (3, "G70", "China"),
        (4, "F70", "EU"),
        (5, "G60", "EU"),
        (6, "G70", "EU"),
    ]:
        for r in range(4, 77, 2):
            value = ws.cell(r, col).value
            if not isinstance(value, (int, float)):
                continue
            label = str(ws.cell(r - 1, 1).value or "Overall product satisfaction")
            loc = f"CN_EU_2025!{openpyxl.utils.get_column_letter(col)}{r}"
            evidence.append(
                Evidence(
                    id=f"studies:mean:{col}:{r}",
                    source_id="studies",
                    domain="customer",
                    location=loc,
                    passage=f"{label}: mean {value}; rating scale and sample size undocumented.",
                    vehicle=vehicle,
                    region=region,
                    period="2025",
                    kind="study_mean",
                    topics=topics_for(label),
                    quantitative={
                        "value": value,
                        "unit": "undocumented rating scale",
                        "n": None,
                    },
                    limitations=s.limitations,
                )
            )
    ws = wb["US_2025"]
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r + 1, 1).value or "").lower().replace(" ", "") != "sampletotal":
            continue
        label = str(ws.cell(r, 1).value)
        for col, vehicle in [(3, "G60"), (4, "G70")]:
            rating_rows = [
                k
                for k in range(r + 2, min(r + 12, ws.max_row + 1))
                if ws.cell(k, 1).value
                in {
                    "I Hate It",
                    "A Failure",
                    "Unsatisfactory",
                    "Satisfactory",
                    "Excellent",
                    "Delightful",
                    "I Love It",
                }
            ][:7]
            distribution = {
                str(ws.cell(k, 1).value): ws.cell(k, col).value for k in rating_rows
            }
            if len(distribution) != 7 or not all(
                isinstance(v, (int, float)) for v in distribution.values()
            ):
                raise ValueError(f"Malformed study distribution at row {r}")
            if abs(sum(distribution.values()) - 1) > 1e-6:
                raise ValueError(f"Unnormalized study distribution at row {r}")
            loc = f"US_2025!{openpyxl.utils.get_column_letter(col)}{r + 1}:{openpyxl.utils.get_column_letter(col)}{max(rating_rows)}"
            evidence.append(
                Evidence(
                    id=f"studies:us:{col}:{r}",
                    source_id="studies",
                    domain="customer",
                    location=loc,
                    passage=f"{label}: normalized response distribution {json.dumps(distribution)}. Respondent count unknown.",
                    vehicle=vehicle,
                    region="United States",
                    period="2025",
                    kind="study_distribution",
                    topics=topics_for(label),
                    quantitative={
                        "distribution": distribution,
                        "unit": "normalized column",
                        "n": None,
                        "sample_total": ws.cell(r + 1, col).value,
                    },
                    limitations=s.limitations,
                )
            )
    wb.close()
    s.coverage = {
        "extracted_records": sum(e.source_id == "studies" for e in evidence),
        "sheets": ["CN_EU_2025", "US_2025"],
    }

    name = "sales_volumes.xlsx"
    s = source(
        name,
        "sales",
        "market",
        source_url=f"https://docs.google.com/spreadsheets/d/{DRIVE_IDS['sales']}/edit",
        limitations=[
            "G60/G68 are aggregated.",
            "2030 values are supplied future figures, not actuals.",
            "Zero volumes do not establish unavailability.",
            "EU is distinct from Europe.",
        ],
    )
    wb = openpyxl.load_workbook(files[name], data_only=True)
    formula_wb = openpyxl.load_workbook(files[name], data_only=False)
    for ws in wb:
        for r in range(2, 6):
            for col, year in [(3, 2024), (4, 2025), (5, 2030)]:
                cell = ws.cell(r, col)
                raw = formula_wb[ws.title].cell(r, col).value
                region = {"US": "United States", "CN": "China"}.get(
                    str(ws.cell(r, 2).value), str(ws.cell(r, 1).value)
                )
                if str(region).lower() in {"europe", "eu"}:
                    region = "EU"
                evidence.append(
                    Evidence(
                        id=f"sales:{ws.title}:{cell.coordinate}",
                        source_id="sales",
                        domain="market",
                        location=f"{ws.title}!{cell.coordinate}",
                        passage=f"{ws.title}, {region}, {year}: {cell.value} vehicles ({'supplied future figure' if year == 2030 else 'supplied historical value'}).",
                        vehicle=ws.title,
                        region=region,
                        period=str(year),
                        kind="provided_forecast" if year == 2030 else "sales_actual",
                        topics=["market-context"],
                        quantitative={
                            "value": cell.value,
                            "unit": "vehicles",
                            "formula": raw
                            if isinstance(raw, str) and raw.startswith("=")
                            else None,
                            "cached_value": cell.value,
                        },
                        limitations=s.limitations,
                    )
                )
    wb.close()
    formula_wb.close()
    s.coverage = {"extracted_records": 36, "sheets": ["F70", "G60_G68", "G70"]}

    for vehicle in FEEDBACK:
        name = f"{vehicle}_OptionList.PDF"
        s = source(
            name,
            f"{vehicle}_equipment",
            "product",
            edition="Juli 2026 / Ausgabe 1 / 28.05.2026",
            limitations=[
                "German edition only; not evidence of Europe-wide availability.",
                "PDF text extraction loses table columns. Equipment/derivative/status/dependencies require visual verification.",
            ],
        )
        from .equipment import matrix_rows

        row_count = 0
        with fitz.open(files[name]) as pdf:
            for i, page in enumerate(pdf):
                passage = page.get_text()
                evidence.append(
                    Evidence(
                        id=f"{s.id}:p{i + 1}",
                        source_id=s.id,
                        domain="product",
                        location=f"PDF page {i + 1}",
                        passage=passage,
                        vehicle=vehicle,
                        region="Germany",
                        published_at="2026-05-28",
                        kind="equipment",
                        topics=topics_for(passage),
                        limitations=s.limitations,
                    )
                )
                rows = matrix_rows(page, vehicle, s.id, i + 1)
                evidence.extend(rows)
                row_count += len(rows)
            s.coverage = {
                "pages": len(pdf),
                "text_pages": len(pdf),
                "verified_equipment_rows": row_count,
                "structured_coverage": "Selected standard-equipment matrices; optional/package tables remain unverified page context",
            }
    name = "BMW-Group-Report-2025-en.pdf"
    s = source(
        name,
        "annual-report",
        "corporate",
        source_url="https://www.bmwgroup.com/content/dam/grpw/websites/bmwgroup_com/ir/downloads/en/2026/bericht/BMW-Group-Report-2025-en.pdf",
        edition="BMW Group Report 2025",
        limitations=[
            "Corporate investment supports alignment, not feature demand, feasibility or sales uplift.",
            "Grouped model families cannot establish single-model volumes.",
            "2026 outlook is not a 2029–2031 forecast.",
            "Results, expectations, targets and announcements are distinct.",
        ],
    )
    pages = [*range(42, 47), *range(53, 58), *range(77, 82), *range(262, 266)]
    with fitz.open(files[name]) as pdf:
        for page in pages:
            text = pdf[page - 1].get_text()
            evidence.append(
                Evidence(
                    id=f"annual-report:p{page}",
                    source_id=s.id,
                    domain="corporate",
                    location=f"PDF page {page} (printed page {page})",
                    passage=text,
                    kind="report_context",
                    context_type="management_expectation"
                    if page >= 262
                    else "reported_result"
                    if 77 <= page <= 81
                    else "mixed",
                    period="2026 outlook" if page >= 262 else "2025 report",
                    topics=topics_for(text),
                    limitations=s.limitations,
                )
            )
        s.coverage = {
            "total_pages": len(pdf),
            "extracted_pages": pages,
            "selection": "User-supplied market/sales and strategy/R&D ranges; financial remainder not extracted",
        }
    warnings.extend(
        [
            "No verified competitor feature comparisons are supplied.",
            "No feedback file contains China observations.",
            "Raw data is private; do not publish source text without authorization.",
        ]
    )
    catalog = Catalog(
        snapshot_hash=snapshot, sources=sources, evidence=evidence, warnings=warnings
    )
    if target:
        target.parent.mkdir(parents=True, exist_ok=True)
        temp = target.with_suffix(".tmp")
        temp.write_text(catalog.model_dump_json())
        temp.replace(target)
    return catalog


def applicable(e: Evidence, scope, context: bool = False) -> bool:
    if e.excluded_reason:
        return False
    if e.vehicle and e.vehicle != scope.vehicle:
        if not (context and e.vehicle == "G60_G68" and scope.vehicle == "G60"):
            return False
    if e.powertrain and e.powertrain != scope.powertrain:
        return False
    if e.country:
        return (
            e.country in EUROPE
            if scope.market == "Europe"
            else e.country == {"China": "CN", "United States": "US"}[scope.market]
        )
    if e.region:
        if e.region == "EU":
            return scope.market == "Europe"
        if context and e.region == "Germany":
            return scope.market == "Europe"
        return e.region == scope.market
    return True


def role_evidence(catalog: Catalog, scope, role: str, ids: list[str]) -> list[Evidence]:
    return [
        e
        for e in catalog.evidence
        if e.id in ids
        and e.domain in {role, "product", "corporate"}
        and applicable(e, scope, True)
    ]

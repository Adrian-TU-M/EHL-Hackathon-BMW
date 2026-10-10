"""Build a local evidence index. Run explicitly; never ingest on an API request."""

import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from backend.schemas.evidence import EvidenceRecord

EUROPE = {"GB", "DE", "FR", "SE", "NO", "DK", "FI", "LV", "ES", "IT", "NL", "AT", "CH", "BE", "IE", "PL"}
POWERTRAINS = {"ICE": "Combustion", "PHEV": "Plug-in hybrid", "BEVE": "Electric", "BEV": "Electric"}


def market_for(country: str) -> str:
    return "Europe" if country in EUROPE else {"US": "United States", "CN": "China"}.get(country, country)


def build_index(data_dir: Path, output: Path, catalog: Path) -> dict:
    now = datetime.now(timezone.utc)
    records: list[EvidenceRecord] = []
    for model in ("F70", "G60", "G70"):
        path = data_dir / f"{model}_feedback_hackathon.xlsx"
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook["Feedback_Explorer"]
            rows = sheet.iter_rows(values_only=True)
            headers = [str(v) for v in next(rows)]
            required = {"ID", "Country", "Engine Type", "Customer Feedback"}
            if not required.issubset(headers):
                raise ValueError(f"Unexpected columns in {path.name}")
            merged: dict[tuple, EvidenceRecord] = {}
            for row_number, values in enumerate(rows, 2):
                row = dict(zip(headers, values))
                comment = str(row.get("Customer Feedback") or "").strip()
                feedback_id = str(row.get("ID") or "").strip()
                if not comment or not feedback_id:
                    continue
                country = str(row["Country"] or "unknown")
                powertrain = POWERTRAINS.get(str(row["Engine Type"]))
                # Preserve distinct comments/scopes sharing an ID; merge only classification duplicates.
                key = (feedback_id, comment, country, powertrain)
                categories = [str(row[k]) for k in headers if k.startswith("Vfc ") and row.get(k)]
                if key in merged:
                    record = merged[key]
                    record.categories = sorted(set(record.categories + categories))
                    record.locator += f", I{row_number}"
                    continue
                digest = hashlib.sha256(json.dumps(key).encode()).hexdigest()[:12]
                merged[key] = EvidenceRecord(
                    id=f"feedback-{model}-{digest}", source_type="customer_feedback",
                    source_title=path.name, source_uri=path.name,
                    locator=f"Feedback_Explorer!I{row_number}", passage=comment,
                    retrieved_at=now, brand="BMW", model=model, market=market_for(country),
                    country=country, powertrain=powertrain, feedback_id=feedback_id,
                    categories=sorted(set(categories)),
                    caveats=["Customer-reported perception, not verified competitor equipment or a representative prevalence estimate."],
                )
            records.extend(merged.values())
        finally:
            workbook.close()

        path = data_dir / f"{model}_OptionList.PDF"
        reader = PdfReader(path)
        for page_number, page in enumerate(reader.pages, 1):
            page_text = page.extract_text(extraction_mode="layout") or ""
            # Keep locators into the extracted page; layout text is NOT a decoded equipment matrix.
            for offset in range(0, len(page_text), 2400):
                passage = page_text[offset:offset + 2400].strip()
                if not passage:
                    continue
                digest = hashlib.sha256(passage.encode()).hexdigest()[:12]
                records.append(EvidenceRecord(
                    id=f"bmw-{model}-p{page_number}-{digest}", source_type="bmw_equipment",
                    source_title=path.name, source_uri=path.name,
                    locator=f"page {page_number}, extracted characters {offset + 1}-{min(offset + 2400, len(page_text))}",
                    passage=passage, retrieved_at=now, brand="BMW", model=model,
                    market="Europe", country="DE",
                    caveats=["German-market option list. Not evidence of US/China availability.",
                             "Extracted table text: trim columns, symbols and footnotes need visual verification before claiming standard/optional/absent.",
                             "Source edition July 2026; future availability is not established."],
                ))
    raw = json.loads(catalog.read_text())
    records.extend(EvidenceRecord.model_validate(item) for item in raw)
    if len({r.id for r in records}) != len(records):
        raise ValueError("Duplicate evidence IDs")
    output.parent.mkdir(parents=True, exist_ok=True)
    payload = {"version": 1, "indexed_at": now.isoformat(), "records": [r.model_dump(mode="json") for r in records]}
    temporary = output.with_suffix(".tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False))
    temporary.replace(output)
    return {"records": len(records), "output": str(output), "competitor_records": len(raw)}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-dir", required=True, type=Path)
    parser.add_argument("--output", type=Path, default=Path("backend/data/evidence.json"))
    parser.add_argument("--catalog", type=Path, default=Path(__file__).parent / "evidence/competitors.json")
    args = parser.parse_args()
    print(json.dumps(build_index(args.data_dir, args.output, args.catalog)))

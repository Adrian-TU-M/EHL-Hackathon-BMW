"""Import selected survey items with exact cell provenance, without inventing scales."""
import hashlib
import math
from openpyxl import load_workbook

ITEMS = {24:('controls','Easy to read and use instruments'),28:('infotainment','Features of information'),30:('infotainment','Easy to use information'),32:('navigation','Integrated navigation'),42:('efficiency','Fuel economy'),74:('connected','Connected services')}

def study_records(path, now):
    if not path.exists(): return []
    digest=hashlib.sha256(path.read_bytes()).hexdigest()
    book=load_workbook(path,read_only=True,data_only=True)
    records=[]
    def add(eid,model,market,tag,locator,passage,caveat):
        records.append(dict(id=eid,source_type='customer_study',source_title='BMW challenge customer study · 2025',source_uri='local:'+path.name,locator=locator+'; sha256:'+digest,passage=passage,retrieved_at=now,brand='BMW study',model=model,market=market,categories=['Feature trends',tag],content_kind='curated_summary',caveats=[caveat,'Sample size, weighting and study provenance are not documented here.','Single-year survey snapshot; not a market trend, feature demand, feature gap or causal sales effect.','All powertrains / unspecified; do not attribute this result to the selected powertrain.']))
    sheet=book['CN_EU_2025']
    for col,model,market in [(2,'G68','China'),(3,'G70','China'),(4,'F70','Europe'),(5,'G60','Europe'),(6,'G70','Europe')]:
        if not str(sheet.cell(3,col).value).strip().endswith(model):raise ValueError('Study model headers changed')
        for row,(tag,prefix) in ITEMS.items():
            label=str(sheet.cell(row-1,1).value).strip()
            # Labels in this workbook can be in the second column; schema must match before importing.
            if not label.startswith(prefix):
                label=str(sheet.cell(row-1,2).value).strip()
            if not label.startswith(prefix):raise ValueError(f'Study item changed: row {row-1}: {label}')
            if sheet.cell(row,1).value!='Mean':raise ValueError('Study statistic label changed')
            value=sheet.cell(row,col).value
            if isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value):raise ValueError('Invalid study mean')
            cell=sheet.cell(row,col).coordinate
            add(f'study-{market}-{model}-{row}',model,market,tag,f'CN_EU_2025!{cell}; item row {row-1}',f'The supplied 2025 {market} {model} study reports a mean of {value:.3f} (rounded) for “{label}”.','Rating scale is not documented here: do not call this a percentage or a score out of ten; do not compare it with US response shares.')
    # US rows are ordinal response shares, not the CN/EU means. Preserve the entire distribution.
    sheet=book['US_2025']
    for col,model in [(3,'G60'),(4,'G70')]:
        if not str(sheet.cell(1,col).value).strip().endswith(model):raise ValueError('US study model headers changed')
        for row,tag in [(775,'navigation'),(829,'controls')]:
            label=str(sheet.cell(row,1).value or sheet.cell(row,2).value).strip()
            expected={775:'If applicable: Navigation system',829:'Ease of reading and using screen interface'}
            if label!=expected[row]:raise ValueError('US study item changed')
            shares=[]
            for r in range(row+2,row+9):
                response=str(sheet.cell(r,1).value).strip()
                value=sheet.cell(r,col).value
                if not isinstance(value,(int,float)) or not math.isfinite(value) or not 0<=value<=1:raise ValueError('Invalid US survey share')
                shares.append((response,value))
            if [name for name,_ in shares]!=['I Hate It','A Failure','Unsatisfactory','Satisfactory','Excellent','Delightful','I Love It']:raise ValueError('US response categories changed')
            if abs(sum(v for _,v in shares)-1)>.01:raise ValueError('US shares do not sum to one')
            text='; '.join(f'{label}: {value*100:.1f}%' for label,value in shares)
            add(f'study-US-{model}-{row}',model,'United States',tag,f'US_2025!{sheet.cell(row,col).column_letter}{row+2}:{sheet.cell(row,col).column_letter}{row+8}; item row {row}',f'The supplied 2025 US {model} study records response shares for “{label}”: {text}.','Ordinal response distribution, not a comparable CN/EU mean. The workbook total of 1 is a normalized share, not a sample of one customer.')
    book.close()
    return records
